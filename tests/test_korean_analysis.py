from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
from openpyxl import load_workbook
from sklearn.feature_extraction.text import CountVectorizer, TfidfVectorizer
from sklearn.metrics import adjusted_rand_score

from survey_insight.edits import apply_edit
from survey_insight.exports import export_excel
from survey_insight.models import TableData, TextDocument
from survey_insight.pipeline import analyze_table, build_cross_analysis
from survey_insight.preprocessing import analyze_morphemes, is_no_opinion, normalize_text, tokenize_keywords
from survey_insight.sentiment import (aggregate_document_sentiment, analyze_text_sentiment,
                                      load_sentiment_lexicon, sentiment_lexicon_info, sentiment_payload)
from survey_insight.topics import _class_tfidf_scores

FIXTURE = json.loads((Path(__file__).parent / "fixtures/korean_analysis.json").read_text(encoding="utf-8"))


class KoreanSentimentTests(unittest.TestCase):
    def test_regression_corpus(self) -> None:
        for text, expected in FIXTURE["sentiment"]:
            with self.subTest(text=text):
                result = analyze_text_sentiment(text)
                self.assertEqual(result.polarity_label, expected)
                for evidence in result.evidence_spans:
                    self.assertEqual(normalize_text(text)[evidence.start:evidence.end], evidence.text)

    def test_negation_stays_inside_clause_and_handles_double_negation(self) -> None:
        cases = [
            ("설명은 좋지만 실습 자료가 없습니다.", "positive"),
            ("좋지 않은 것은 아닙니다.", "positive"),
            ("불편함이 없어요.", "positive"),
            ("전혀 유용하지 않아요.", "negative"),
            ("내용이 좋습니다. 일정은 만족스럽지 않습니다.", "mixed"),
            ("불친절하지 않아요.", "positive"),
            ("조금 불편하지만 아주 유익합니다.", "mixed"),
            ("별로 좋지 않습니다.", "negative"),
            ("별로 나쁘지 않습니다.", "positive"),
            ("시설이 별로예요.", "negative"),
            ("휴가를 눈치 없이 쓸 수 있어 좋습니다.", "positive"),
            ("월급이 적어 생활이 빠듯합니다.", "negative"),
            ("야근이 잦아 지칩니다.", "negative"),
        ]
        for text, expected in cases:
            with self.subTest(text=text):
                self.assertEqual(analyze_text_sentiment(text).polarity_label, expected)

    def test_quoted_judgment_requires_review(self) -> None:
        result = analyze_text_sentiment("서비스가 좋다는 말에 동의하지 않습니다.")
        self.assertTrue(result.needs_review)

    def test_aggregation_reuses_saved_document_analysis(self) -> None:
        document = TextDocument("saved", 1, "text", "교육이 유익합니다", "교육이 유익합니다", {})
        document.sentiment = sentiment_payload(analyze_text_sentiment(document.redacted_text))
        with patch("survey_insight.sentiment.analyze_text_sentiment", side_effect=AssertionError("must reuse saved result")):
            result = aggregate_document_sentiment([document, document])
        self.assertEqual(result.count, 2)
        self.assertEqual(result.label_counts["positive"], 2)

    def test_request_and_no_evidence_are_distinguishable(self) -> None:
        request = analyze_text_sentiment("온라인 교육 과정 개선이 필요합니다.")
        no_evidence = analyze_text_sentiment("신청서는 내일 제출합니다.")
        self.assertEqual(request.polarity_label, "neutral")
        self.assertEqual(request.evidence_status, "request_only")
        self.assertEqual(no_evidence.evidence_status, "no_match")
        self.assertTrue(no_evidence.needs_review)
        self.assertEqual(analyze_text_sentiment("시급하지 않습니다.").urgency_score, 0)

    def test_dictionary_is_bundled_and_versioned(self) -> None:
        info = sentiment_lexicon_info()
        self.assertGreaterEqual(len(load_sentiment_lexicon()), 150)
        self.assertEqual(info["version"], "1.0.0")
        self.assertEqual(len(info["sha256"]), 64)

    def test_unavailable_morphology_is_explicit_degraded_path(self) -> None:
        analyze_text_sentiment.cache_clear()
        with patch("survey_insight.sentiment.tokenizer_source", return_value="regex_keyword_fallback"):
            result = analyze_text_sentiment("이 교육은 불친절합니다.")
        analyze_text_sentiment.cache_clear()
        self.assertEqual(result.polarity_label, "negative")
        self.assertEqual(result.source, "korean_surface_lexicon_fallback")
        self.assertTrue(result.needs_review)

    def test_normalization_and_topic_pos(self) -> None:
        self.assertEqual(normalize_text("ＡＩ\u200b 교육"), "AI 교육")
        self.assertTrue(is_no_opinion("없습니다."))
        self.assertTrue(is_no_opinion(float("nan")))
        self.assertFalse(is_no_opinion("주차장이 없습니다."))
        tokens = tokenize_keywords("AI 앱의 신청 방법을 안내해주세요. [EMAIL] https://example.com")
        self.assertIn("ai", tokens)
        self.assertIn("앱", tokens)
        self.assertNotIn("email", tokens)
        self.assertNotIn("example", tokens)
        self.assertIn("않", [t.form for t in analyze_morphemes("만족하지 않습니다")])


class KoreanTopicIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.rows = [{"response": text, "gold": gold} for gold, texts in FIXTURE["topics"].items() for text in texts]
        cls.package = analyze_table(TableData(["response", "gold"], cls.rows), text_column="response", group_columns=[])

    def test_recovers_subjects_across_mixed_polarities(self) -> None:
        by_id = {a.document_id: a.topic_id for a in self.package.assignments}
        predicted = [by_id[d.id] for d in self.package.documents]
        self.assertEqual(len(self.package.selected_topics), 3)
        self.assertGreaterEqual(adjusted_rand_score([r["gold"] for r in self.rows], predicted), 0.9)

    def test_document_sentiment_and_topic_counts_reconcile(self) -> None:
        p = self.package
        self.assertEqual(sum(p.recommendation.sentiment_summary["label_counts"].values()), len(self.rows))
        self.assertTrue(all(d.sentiment for d in p.documents))
        self.assertTrue(all(d.topic_tokens for d in p.documents))
        for topic in p.selected_topics:
            self.assertEqual(sum(topic.sentiment_distribution.values()), topic.count)
        for assignment in p.assignments:
            self.assertAlmostEqual(sum(assignment.membership_scores.values()), 1.0, places=4)
            self.assertNotEqual(assignment.score_method, "unspecified")

    def test_export_contains_redacted_evidence_and_distribution(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            book = load_workbook(export_excel(self.package, Path(tmp) / "report.xlsx"))
        self.assertIn("감정분포", book.sheetnames)
        headers = [cell.value for cell in book["원자료_토픽배정"][1]]
        for name in ["sentiment_label", "sentiment_evidence", "score_method", "topic_tokens"]:
            self.assertIn(name, headers)
        self.assertNotIn("original_text", headers)

    def test_repeating_three_subjects_does_not_force_more_topics(self) -> None:
        rows = self.rows * 3
        p = analyze_table(TableData(["response", "gold"], rows), text_column="response", group_columns=[])
        by_id = {a.document_id: a.topic_id for a in p.assignments}
        self.assertEqual(len(p.selected_topics), 3)
        self.assertGreaterEqual(adjusted_rand_score([r["gold"] for r in rows], [by_id[d.id] for d in p.documents]), 0.9)

    def test_empty_topic_vocabulary_preserves_sentiment(self) -> None:
        for texts in [["좋습니다", "만족합니다"], ["ㅋㅋㅋ", "!", "a"], [""], ["교육"]]:
            with self.subTest(texts=texts):
                p = analyze_table(TableData(["response"], [{"response": s} for s in texts]), text_column="response", group_columns=[])
                self.assertEqual(len(p.documents), len(texts))
                self.assertLessEqual(len(p.selected_topics), 1)
                if texts[0] == "좋습니다":
                    self.assertEqual(p.documents[0].sentiment["polarity_label"], "positive")
                    self.assertEqual(p.assignments[0].topic_id, "OUTLIER")

    def test_class_tfidf_does_not_reuse_tfidf_as_counts(self) -> None:
        texts = ["alpha alpha common", "alpha common", "beta beta common", "beta common"]
        docs = [TextDocument(str(i), i, "text", text, text, {}) for i, text in enumerate(texts)]
        tfidf = TfidfVectorizer()
        matrix = tfidf.fit_transform(texts)
        counter = CountVectorizer(vocabulary=tfidf.vocabulary_)
        counts = counter.fit_transform(texts)
        labels = np.asarray([0, 0, 1, 1])
        first = _class_tfidf_scores(docs, labels, tfidf, matrix)
        second = _class_tfidf_scores(docs, labels, counter, counts)
        np.testing.assert_allclose(first[0], second[0])

    def test_merging_topics_recalculates_sentiment(self) -> None:
        from copy import deepcopy
        package = deepcopy(self.package)
        ids = [topic.topic_id for topic in package.selected_topics[:2]]
        total = sum(topic.count for topic in package.selected_topics[:2])
        apply_edit(package, "merge_topics", {"source_topic_ids": ids, "target_label": "병합 검토"})
        topic = next(topic for topic in package.selected_topics if topic.topic_id == ids[0])
        self.assertEqual(topic.count, total)
        self.assertEqual(sum(topic.sentiment_distribution.values()), total)
        self.assertEqual(topic.sentiment_method, "korean_morph_lexicon_v1")

    def test_small_group_sentiment_cells_are_suppressed(self) -> None:
        from copy import deepcopy
        package = deepcopy(self.package)
        for index, document in enumerate(package.documents):
            document.metadata["group"] = "small" if index == 0 else "large"
        rows = build_cross_analysis(package.documents, package.assignments, package.selected_topics, ["group"])
        small = next(row for row in rows if row["group_value"] == "small")
        self.assertTrue(small["suppressed"])
        self.assertIsNone(small["display_count"])
        self.assertIsNone(small["sentiment_counts"])
        for row in rows:
            if not row["suppressed"]:
                self.assertEqual(sum(row["sentiment_counts"].values()), row["count"])
