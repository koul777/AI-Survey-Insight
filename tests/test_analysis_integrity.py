from __future__ import annotations

from pathlib import Path
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook

from survey_insight.edits import apply_edit
from survey_insight.exports import export_excel, export_powerpoint, export_word
from survey_insight.models import (
    AnalysisPackage, CandidateSolution, DatasetProfile, Recommendation,
    TextDocument, Topic, TopicAssignment, to_jsonable,
)
from survey_insight.storage import StorageRepository


def _package() -> AnalysisPackage:
    documents = []
    assignments = []
    topics = []
    labels = ("positive", "negative", "neutral")
    for topic_index in range(3):
        topic_id = f"T0{topic_index + 1}"
        texts = [f"주제{topic_index + 1} 의견{index + 1}" for index in range(3)]
        for index, text in enumerate(texts):
            document_id = f"doc_{topic_index}_{index}"
            documents.append(TextDocument(
                document_id, len(documents), "response", text, text, {"team": "small"},
                sentiment={"polarity_label": labels[topic_index], "polarity_score": (0.6, -0.6, 0)[topic_index],
                           "urgency_score": 0.0, "evidence_terms": [], "evidence_status": "matched",
                           "source": "fixture_lexicon"},
            ))
            membership = {"T01": 0.3, "T02": 0.3, "T03": 0.4} if topic_index == 2 else {
                key: 0.6 if key == topic_id else 0.2 for key in ("T01", "T02", "T03")
            }
            assignments.append(TopicAssignment(
                document_id, topic_id, membership[topic_id], False, "model",
                membership_scores=membership, score_method="normalized_nmf_membership",
            ))
        topics.append(Topic(topic_id, f"주제{topic_index + 1}", "설명", 3, 1 / 3, ["교육", "연수"], texts,
                            sentiment_distribution={label: 3 if label == labels[topic_index] else 0 for label in labels},
                            sentiment_evidence_coverage=1.0))
    candidate = CandidateSolution("candidate", 3, {"quality_gate_passed": True}, {"coverage": 1.0}, 0.8,
                                  topics=topics, assignments=assignments)
    recommendation = Recommendation("run_integrity", "test", 9, (1, 3), 1, candidate, None, None, [candidate])
    return AnalysisPackage({"name": "Integrity"}, DatasetProfile("dataset_integrity", None, 9, 2, 0, [], []),
                           "response", ["team"], documents, recommendation, "candidate", topics, assignments, [])


class AnalysisIntegrityTests(unittest.TestCase):
    def test_legacy_chat_units_are_recovered_on_read_without_rewriting_storage(self) -> None:
        package = _package()
        package.project["chat_preprocessing"] = {"retained_messages": 27, "context_segments": 9}
        package.recommendation.sentiment_summary = {"analyzed_count": 9}
        with tempfile.TemporaryDirectory() as tmp:
            repo = StorageRepository(Path(tmp) / "legacy-chat.db")
            original = repo.save_run(package=package)
            restored = original.to_analysis_package()
            compact = repo.get_run_recommendation_json("run_integrity")
            self.assertEqual(restored.recommendation.sentiment_summary["analysis_unit"], "context_segment")
            self.assertEqual(compact["sentiment_summary"], restored.recommendation.sentiment_summary)
            self.assertEqual(compact["sentiment_summary"]["source_message_count"], 27)
            self.assertNotIn("_source_row_indices", restored.documents[0].metadata)
            reloaded = repo.get_run("run_integrity")
            self.assertEqual(reloaded.package_json, original.package_json)
            self.assertEqual(reloaded.updated_at, original.updated_at)
            self.assertNotIn("analysis_unit", original.package_json["recommendation"]["sentiment_summary"])

    def test_legacy_chat_exports_do_not_invent_original_row_numbers(self) -> None:
        package = _package()
        package.project["chat_preprocessing"] = {"retained_messages": 27, "context_segments": 9}
        package.cross_analysis = [{"group_column": "author", "group_value": "LEGACY_AUTHOR_GROUP",
                                   "topic_id": "T01", "count": 9, "display_count": 9,
                                   "share": 1.0, "suppressed": False}]
        with tempfile.TemporaryDirectory() as tmp:
            workbook = load_workbook(export_excel(package, Path(tmp) / "legacy-chat.xlsx"))
            rows = list(workbook["원자료_토픽배정"].values)
            self.assertEqual(rows[1][rows[0].index("source_row_indices")], "[]")
            self.assertEqual(rows[1][rows[0].index("source_row_mapping_status")], "not_recorded")
            summary = list(workbook["감정분포"].values)
            self.assertIn("대화 맥락", summary[1][0])
            self.assertTrue(any(row[0] == "source_message_count" and row[1] == 27 for row in summary))
            self.assertNotIn("LEGACY_AUTHOR_GROUP", str(list(workbook["교차분석"].values)))
            target = export_word(package, Path(tmp) / "legacy-chat.docx")
            with zipfile.ZipFile(target) as archive:
                text = archive.read("word/document.xml").decode("utf-8")
                self.assertIn("원본 행 위치가 기록되지 않았습니다", text)
                self.assertNotIn("LEGACY_AUTHOR_GROUP", text)
            self.assertEqual(package.cross_analysis[0]["group_value"], "LEGACY_AUTHOR_GROUP")

    def test_legacy_group_shares_cannot_reveal_complementary_small_counts(self) -> None:
        package = _package()
        package.cross_analysis = [
            {"group_column": "team", "group_value": "A", "topic_id": "T01", "topic_label": "다수",
             "count": 5, "display_count": 5, "share": 0.8333, "suppressed": False},
            {"group_column": "team", "group_value": "A", "topic_id": "T02", "topic_label": "소수",
             "count": 1, "display_count": None, "share": 0.1667, "suppressed": True},
        ]
        with tempfile.TemporaryDirectory() as tmp:
            workbook = load_workbook(export_excel(package, Path(tmp) / "groups.xlsx"))
            rows = list(workbook["교차분석"].values)
            self.assertEqual(rows[1][4], 5)
            self.assertIsNone(rows[1][5])
            self.assertIsNone(rows[2][5])
            target = export_word(package, Path(tmp) / "groups.docx")
            with zipfile.ZipFile(target) as archive:
                self.assertNotIn("83.3%", archive.read("word/document.xml").decode("utf-8"))
            self.assertEqual(package.cross_analysis[0]["share"], 0.8333)

    def test_chat_exports_keep_message_provenance_and_analysis_units(self) -> None:
        package = _package()
        package.recommendation.sentiment_summary = {"analysis_unit": "context_segment", "source_message_count": 17}
        package.documents[0].metadata.update({"_source_row_indices": [2, 4, 7], "_source_message_count": 3})
        with tempfile.TemporaryDirectory() as tmp:
            workbook = load_workbook(export_excel(package, Path(tmp) / "chat.xlsx"))
            rows = list(workbook["원자료_토픽배정"].values)
            self.assertEqual(rows[1][rows[0].index("source_row_indices")], "[2, 4, 7]")
            self.assertEqual(rows[1][rows[0].index("source_message_count")], 3)
            summary_rows = list(workbook["감정분포"].values)
            self.assertIn("대화 맥락", summary_rows[1][0])
            self.assertTrue(any(row[0] == "source_message_count" and row[1] == 17 for row in summary_rows))
            for exporter, suffix in ((export_word, "docx"), (export_powerpoint, "pptx")):
                target = exporter(package, Path(tmp) / f"chat.{suffix}")
                with zipfile.ZipFile(target) as archive:
                    xml = " ".join(archive.read(name).decode("utf-8") for name in archive.namelist() if name.endswith(".xml"))
                    self.assertIn("서로 다른 발화자", xml)
                    self.assertIn("대화 맥락", xml)

    def test_exclusion_cleans_memberships_and_records_why_documents_are_unassigned(self) -> None:
        package = _package()
        apply_edit(package, "exclude_topic", {"topic_id": "T01"})
        self.assertEqual(len(package.assignments), 6)
        for assignment in package.assignments:
            self.assertNotIn("T01", assignment.membership_scores)
            self.assertAlmostEqual(sum(assignment.membership_scores.values()), 1.0, places=5)
            self.assertEqual(assignment.probability, assignment.membership_scores[assignment.topic_id])
            self.assertEqual(assignment.score_method, "membership_conditional_on_retained_topics")
        self.assertTrue(all(doc.topic_exclusion_reason == "user_excluded_topic" for doc in package.documents[:3]))
        self.assertEqual([topic.share for topic in package.selected_topics], [0.5, 0.5])
        self.assertTrue(all("50.0%" in topic.plain_language_summary for topic in package.selected_topics))
        self.assertTrue(all("현재 남은 토픽 배정 중" in topic.plain_language_summary for topic in package.selected_topics))

    def test_missing_topic_exclusion_fails_without_mutation(self) -> None:
        package = _package()
        before = to_jsonable(package)
        with self.assertRaises(ValueError):
            apply_edit(package, "exclude_topic", {"topic_id": "does-not-exist"})
        self.assertEqual(to_jsonable(package), before)

    def test_merge_preserves_source_examples_sentiment_and_combines_memberships(self) -> None:
        package = _package()
        apply_edit(package, "merge_topics", {"source_topic_ids": ["T01", "T02"], "target_label": "병합"})
        topic = package.selected_topics[0]
        self.assertEqual(topic.count, 6)
        self.assertEqual(topic.sentiment_distribution, {"positive": 3, "neutral": 0, "negative": 3, "mixed": 0})
        self.assertEqual(topic.sentiment_method, "fixture_lexicon")
        self.assertIn("주제1 의견1", topic.representative_responses)
        self.assertIn("주제2 의견1", topic.representative_responses)
        self.assertTrue(topic.sentiment_plain_language)
        for assignment in package.assignments:
            self.assertNotIn("T02", assignment.membership_scores)
            self.assertAlmostEqual(sum(assignment.membership_scores.values()), 1.0)
            self.assertTrue(assignment.needs_review)
        self.assertEqual(package.assignments[-1].topic_id, "T03")
        self.assertEqual(package.assignments[-1].probability, 0.4)

    def test_structural_edit_keeps_diagnostics_scoped_to_original_model(self) -> None:
        package = _package()
        apply_edit(package, "merge_topics", {"source_topic_ids": ["T01", "T02"], "target_label": "병합"})
        apply_edit(package, "exclude_topic", {"topic_id": "T03"})
        candidate = package.recommendation.recommended
        self.assertEqual(candidate.topic_count, 1)
        self.assertEqual(candidate.params["original_model_diagnostics"]["topic_count"], 3)
        self.assertEqual(candidate.params["original_model_diagnostics"]["score"], 0.8)
        self.assertEqual(candidate.params["diagnostics_status"], "invalidated_by_user_edit")
        self.assertFalse(candidate.params["quality_gate_passed"])
        self.assertIn("편집 전", package.recommendation.plain_language_summary)

    def test_invalid_merge_label_does_not_mutate_package(self) -> None:
        package = _package()
        before = to_jsonable(package)
        with self.assertRaises(ValueError):
            apply_edit(package, "merge_topics", {"source_topic_ids": ["T01", "T02"], "target_label": " "})
        self.assertEqual(to_jsonable(package), before)

    def test_representative_edit_masks_pii_and_records_actual_change(self) -> None:
        package = _package()
        edit = apply_edit(package, "replace_representative", {
            "topic_id": "T01", "response": "연락은 author@example.com 또는 010-1234-5678",
        })
        response = package.selected_topics[0].representative_responses[0]
        self.assertIn("[EMAIL]", response)
        self.assertIn("[PHONE]", response)
        self.assertNotEqual(edit.before["topics"][0]["representative_responses"],
                            edit.after["topics"][0]["representative_responses"])

    def test_exported_edit_history_cannot_reintroduce_old_pii(self) -> None:
        package = _package()
        package.selected_topics[0].representative_responses = ["old@example.com 문의"]
        apply_edit(package, "replace_representative", {"topic_id": "T01", "response": "대체 의견"})
        with tempfile.TemporaryDirectory() as tmp:
            target = export_excel(package, Path(tmp) / "history.xlsx")
            with zipfile.ZipFile(target) as archive:
                for name in archive.namelist():
                    if name.endswith(".xml"):
                        self.assertNotIn(b"old@example.com", archive.read(name))

    def test_excel_stores_formula_like_input_as_text_in_every_sheet(self) -> None:
        package = _package()
        package.documents[0].redacted_text = '=HYPERLINK("https://example.com", "click")'
        package.documents[1].redacted_text = "#N/A"
        package.documents[2].redacted_text = "safe\x00\x01text"
        package.documents[0].metadata["team"] = "=1+1"
        package.selected_topics[0].label = "=1+1"
        package.selected_topics[0].representative_responses = ["=SUM(1,2)"]
        with tempfile.TemporaryDirectory() as tmp:
            target = export_excel(package, Path(tmp) / "report.xlsx")
            workbook = load_workbook(target)
            self.assertEqual(workbook["원자료_토픽배정"]["B2"].value, package.documents[0].redacted_text)
            self.assertEqual(workbook["원자료_토픽배정"]["B3"].data_type, "s")
            self.assertEqual(workbook["원자료_토픽배정"]["B4"].value, "safetext")
            self.assertEqual(workbook["원자료_토픽배정"]["A2"].data_type, "n")
            self.assertTrue(all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row))
            with zipfile.ZipFile(target) as archive:
                for name in archive.namelist():
                    if name.startswith("xl/worksheets/"):
                        self.assertNotIn(b"<f>", archive.read(name))

    def test_exports_suppress_old_stored_small_cell_shares_and_sentiment(self) -> None:
        package = _package()
        package.cross_analysis = [{"group_column": "team", "group_value": "small", "topic_id": "T01",
                                   "topic_label": "주제1", "count": 1, "display_count": 1, "share": 0.1234,
                                   "suppressed": True, "sentiment_counts": {"positive": 1}}]
        with tempfile.TemporaryDirectory() as tmp:
            target = export_excel(package, Path(tmp) / "report.xlsx")
            row = list(load_workbook(target)["교차분석"].values)[1]
            self.assertIsNone(row[4])
            self.assertIsNone(row[5])
            self.assertTrue(all(value is None for value in row[7:]))
            target = export_word(package, Path(tmp) / "report.docx")
            with zipfile.ZipFile(target) as archive:
                text = archive.read("word/document.xml").decode("utf-8")
                self.assertNotIn("12.3%", text)
                ET.fromstring(text)

    def test_old_representatives_are_redacted_in_all_report_formats(self) -> None:
        package = _package()
        package.selected_topics[0].representative_responses = ["author@example.com 연락 바랍니다"]
        with tempfile.TemporaryDirectory() as tmp:
            for exporter, suffix in ((export_excel, "xlsx"), (export_word, "docx"), (export_powerpoint, "pptx")):
                target = exporter(package, Path(tmp) / f"report.{suffix}")
                with zipfile.ZipFile(target) as archive:
                    for name in archive.namelist():
                        if name.endswith(".xml"):
                            self.assertNotIn(b"author@example.com", archive.read(name))

    def test_legacy_saved_package_defaults_and_exports_survive_round_trip(self) -> None:
        data = to_jsonable(_package())
        data["recommendation"].pop("sentiment_summary")
        data["recommendation"].pop("analysis_diagnostics")
        for document in data["documents"]:
            for key in ("topic_tokens", "sentiment", "topic_exclusion_reason"):
                document.pop(key)
        for topic in data["selected_topics"]:
            for key in ("sentiment_distribution", "sentiment_evidence_coverage", "sentiment_distribution_method"):
                topic.pop(key)
        for assignment in data["assignments"]:
            for key in ("membership_scores", "score_method", "needs_review"):
                assignment.pop(key)
        data["future_extension"] = "ignored"
        with tempfile.TemporaryDirectory() as tmp:
            repo = StorageRepository(Path(tmp) / "legacy.db")
            repo.save_run(package=data)
            restored = StorageRepository(Path(tmp) / "legacy.db").get_run("run_integrity").to_analysis_package()
            self.assertEqual(restored.documents[0].sentiment, {})
            self.assertEqual(restored.assignments[0].membership_scores, {})
            target = export_excel(restored, Path(tmp) / "legacy.xlsx")
            summary = list(load_workbook(target)["감정분포"].values)
            self.assertEqual(summary[2][6], 0)
            self.assertEqual(summary[2][7], "not_available_in_stored_result")


if __name__ == "__main__":
    unittest.main()
