from __future__ import annotations

import tempfile
import unittest
import warnings
import zipfile
from pathlib import Path
from unittest.mock import patch

import numpy as np
from openpyxl import load_workbook
from pptx import Presentation
from sklearn.exceptions import ConvergenceWarning
from sklearn.feature_extraction.text import CountVectorizer

from survey_insight.cli import create_demo_workbook
from survey_insight.edits import apply_edit
from survey_insight.embeddings import EmbeddingResult
from survey_insight.exports import export_excel, export_powerpoint, export_word
from survey_insight.models import Topic
from survey_insight.pipeline import analyze_file
from survey_insight.preprocessing import (
    detect_privacy_review_flags,
    is_no_opinion,
    mask_pii,
    tokenize_keywords,
    tokenizer_source,
)
from survey_insight.topics import (
    _cluster,
    _coherence,
    _count_vectorize,
    _fit_nmf,
    _lda_input_diagnostics,
    _silhouette,
    topic_policy,
)


class PipelineTests(unittest.TestCase):
    def test_nonconverged_nmf_candidate_is_rejected(self) -> None:
        class NonConvergingNMF:
            def __init__(self, **_: object) -> None:
                self.components_ = np.ones((2, 3), dtype=float)

            def fit_transform(self, matrix: np.ndarray) -> np.ndarray:
                warnings.warn("iteration limit", ConvergenceWarning)
                return np.asarray([[1.0, 0.0], [0.0, 1.0], [1.0, 0.0], [0.0, 1.0]])

        with patch("survey_insight.topics.NMF", NonConvergingNMF):
            result = _fit_nmf(np.ones((4, 3), dtype=float), 2, 42)

        self.assertIsNone(result)

    def test_prd_like_fixture_flow(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            workbook_path = tmp_path / "demo.xlsx"
            create_demo_workbook(workbook_path)
            package = analyze_file(workbook_path, project_name="Fixture")
            self.assertEqual(package.dataset_profile.row_count, 94)
            self.assertEqual(package.dataset_profile.column_count, 63)
            self.assertEqual(package.text_column, "15. 인사제도 개선을 위한 기타 건의 사항 [의견]")
            text_profile = next(
                col for col in package.dataset_profile.columns if col.column_name == package.text_column
            )
            self.assertGreaterEqual(text_profile.text_score, 0.70)
            self.assertEqual(package.recommendation.valid_response_count, 16)
            self.assertEqual(package.recommendation.mode, "소표본 모드")
            self.assertGreaterEqual(package.recommendation.recommended.topic_count, 2)
            self.assertLessEqual(package.recommendation.recommended.topic_count, 4)
            self.assertIsNotNone(package.recommendation.detailed)
            engines = {candidate.params["engine"] for candidate in package.recommendation.candidates}
            self.assertIn("kmeans", engines)
            self.assertIn("nmf", engines)
            self.assertIn("lda", engines)
            self.assertTrue(package.recommendation.recommended.params["quality_gate_passed"])
            resampling_status = package.recommendation.recommended.params["resampling_status"]
            self.assertIn(
                resampling_status,
                {"evaluated", "insufficient_successful_repeats", "unsupported_for_density_candidate"},
            )
            if resampling_status == "evaluated":
                self.assertGreaterEqual(package.recommendation.recommended.metrics["resampling_stability"], 0.0)
                self.assertLessEqual(package.recommendation.recommended.metrics["resampling_stability"], 1.0)
            lda_candidate = next(
                candidate for candidate in package.recommendation.candidates if candidate.params["engine"] == "lda"
            )
            self.assertEqual(lda_candidate.params["feature_space"], "term_count")
            self.assertEqual(lda_candidate.params["topic_terms"], "lda_components")
            self.assertTrue(all(topic.keywords for topic in lda_candidate.topics))
            self.assertTrue(any(0.0 < assignment.probability < 1.0 for assignment in lda_candidate.assignments))
            self.assertTrue(package.recommendation.plain_language_summary)
            self.assertIn("유효 자유응답", package.recommendation.plain_language_summary)
            self.assertIn("함께 비교", package.recommendation.topic_count_explanation)
            self.assertTrue(package.recommendation.methodology_plain_language)
            self.assertTrue(package.recommendation.quality_warnings)
            self.assertIn("tokenizer", package.recommendation.recommended.params)
            self.assertEqual(package.selected_topics[0].sentiment_method, "rule_lexicon_fallback")
            self.assertTrue(package.selected_topics[0].plain_language_summary)
            self.assertTrue(package.selected_topics[0].suggested_action)
            self.assertTrue(package.selected_topics[0].sentiment_plain_language)
            package.documents[0].original_text = "연락처는 privacy@example.com 입니다."
            package.documents[0].redacted_text = "연락처는 [EMAIL] 입니다."
            package.documents[0].privacy_review_flags = ["explicit_name_context"]
            excel_path = export_excel(package, tmp_path / "report.xlsx")
            word_path = export_word(package, tmp_path / "report.docx")
            ppt_path = export_powerpoint(package, tmp_path / "report.pptx")
            self.assertTrue(excel_path.exists())
            self.assertTrue(word_path.exists())
            self.assertTrue(ppt_path.exists())
            workbook = load_workbook(excel_path)
            self.assertIn("토픽요약", workbook.sheetnames)
            assigned_sheet = workbook["원자료_토픽배정"]
            assigned_headers = [cell.value for cell in assigned_sheet[1]]
            self.assertNotIn("original_text", assigned_headers)
            self.assertIn("redacted_text", assigned_headers)
            self.assertIn("privacy_review_flags", assigned_headers)
            assigned_values = " ".join(str(cell.value or "") for row in assigned_sheet.iter_rows() for cell in row)
            self.assertNotIn("privacy@example.com", assigned_values)
            self.assertIn("[EMAIL]", assigned_values)
            self.assertIn("explicit_name_context", assigned_values)
            model_settings = " ".join(
                str(cell.value or "")
                for row in workbook["모델설정"].iter_rows()
                for cell in row
            )
            for term in ["군집 품질 종합점수", "토픽 해석 가능성 점수", "군집 분리도", "분석 포함률"]:
                self.assertIn(term, model_settings)
            self.assertIn("human_review.status", model_settings)
            self.assertIn("부분표본 일치도", model_settings)
            with zipfile.ZipFile(word_path) as zf:
                self.assertIn("word/document.xml", zf.namelist())
                word_xml = zf.read("word/document.xml").decode("utf-8")
            self.assertIn("권장 토픽 수", word_xml)
            self.assertIn("군집 품질 종합점수", word_xml)
            self.assertNotIn("제품 추천 토픽 수", word_xml)
            with zipfile.ZipFile(ppt_path) as zf:
                self.assertIn("ppt/presentation.xml", zf.namelist())
            presentation = Presentation(ppt_path)
            ppt_text = "\n".join(
                shape.text
                for slide in presentation.slides
                for shape in slide.shapes
                if hasattr(shape, "text_frame")
            )
            self.assertIn("권장 토픽 수", ppt_text)
            self.assertIn("군집 품질 종합점수", ppt_text)
            self.assertIn("긴급성 판정이 아닙니다", ppt_text)

    def test_preprocessing_privacy_and_no_opinion(self) -> None:
        self.assertTrue(is_no_opinion("해당 없음"))
        self.assertTrue(is_no_opinion("."))
        masked, found = mask_pii("연락처는 010-1234-5678 test@example.com 입니다.")
        self.assertIn("phone", found)
        self.assertIn("email", found)
        self.assertIn("[PHONE]", masked)
        self.assertIn("[EMAIL]", masked)
        employee_masked, employee_found = mask_pii("담당 사번 AB-20260001의 확인이 필요합니다.")
        self.assertIn("employee_id", employee_found)
        self.assertIn("[EMPLOYEE_ID]", employee_masked)
        self.assertNotIn("AB-20260001", employee_masked)
        ordinary_text = "사번 제도와 직번 관리 절차를 개선해 주세요."
        self.assertEqual(mask_pii(ordinary_text), (ordinary_text, []))
        self.assertEqual(
            detect_privacy_review_flags("성명: 홍길동, 주소: 서울, 계좌번호를 적었습니다"),
            ["explicit_name_context", "explicit_address_context", "financial_account_context"],
        )

    def test_kiwi_noun_tokenizer_extracts_korean_nouns(self) -> None:
        self.assertEqual(tokenizer_source(), "kiwi_noun_extractor")
        tokens = tokenize_keywords("평가 기준이 어렵습니다. 인사제도 개선이 필요합니다.")
        self.assertIn("평가", tokens)
        self.assertIn("기준", tokens)
        self.assertIn("인사", tokens)
        self.assertIn("제도", tokens)
        self.assertNotIn("어렵습니다", tokens)

    def test_invalid_silhouette_candidate_gets_zero_separation_score(self) -> None:
        matrix = np.asarray([[1.0, 0.0], [0.9, 0.1], [0.0, 1.0]])
        labels = np.asarray([0, 0, 0])
        self.assertEqual(_silhouette(matrix, labels), 0.0)

    def test_lda_feature_matrix_uses_integer_term_counts(self) -> None:
        _, matrix = _count_vectorize(
            [
                "평가 보상 기준 개선",
                "평가 기준 설명 필요",
                "교육 과정 신청 편의",
                "교육 일정 안내 필요",
            ]
        )
        self.assertTrue(np.all(np.equal(matrix.data, np.floor(matrix.data))))
        self.assertGreater(matrix.sum(), 0)

    def test_short_lda_input_is_flagged_as_insufficient_evidence(self) -> None:
        _, matrix = _count_vectorize(["급여", "평가", "승진", "교육"])
        diagnostics = _lda_input_diagnostics(matrix)
        self.assertFalse(diagnostics["input_evidence_sufficient"])
        self.assertEqual(diagnostics["input_short_response_share"], 1.0)

    def test_collapsed_kmeans_candidate_is_rejected(self) -> None:
        identical = np.ones((6, 3), dtype=float)
        self.assertIsNone(_cluster(identical, 3, seed=42))

    def test_silhouette_excludes_dbscan_outliers(self) -> None:
        matrix = np.asarray(
            [
                [1.0, 0.0],
                [0.95, 0.05],
                [0.0, 1.0],
                [0.05, 0.95],
                [0.5, 0.5],
            ]
        )
        labels = np.asarray([0, 0, 1, 1, -1])
        self.assertGreater(_silhouette(matrix, labels), 0.8)

    def test_topic_coherence_uses_document_cooccurrence(self) -> None:
        vectorizer = CountVectorizer()
        matrix = vectorizer.fit_transform(
            [
                "pay reward policy",
                "pay reward review",
                "pay reward fairness",
                "training schedule course",
                "training schedule registration",
            ]
        )
        coherent = Topic("T01", "보상", "", 3, 0.6, ["pay", "reward"], ["pay reward policy"])
        incoherent = Topic("T02", "혼합", "", 2, 0.4, ["pay", "training"], ["training schedule course"])
        self.assertGreater(
            _coherence(vectorizer, matrix, [coherent]),
            _coherence(vectorizer, matrix, [incoherent]),
        )

    def test_small_sample_policy_is_explicit_and_bounded(self) -> None:
        mode, topic_range, minimum_size, warnings = topic_policy(12)
        self.assertEqual(mode, "초소표본 모드")
        self.assertEqual(topic_range, (1, 3))
        self.assertEqual(minimum_size, 2)
        self.assertTrue(any("테마 코딩" in warning for warning in warnings))

    def test_same_input_and_seed_produce_same_recommendation(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "demo.xlsx"
            create_demo_workbook(workbook_path)
            first = analyze_file(workbook_path, seed=2026)
            second = analyze_file(workbook_path, seed=2026)

        first_signature = (
            first.recommendation.recommended.params["engine"],
            first.recommendation.recommended.topic_count,
            first.recommendation.recommended.score,
            [(topic.count, topic.keywords) for topic in first.selected_topics],
        )
        second_signature = (
            second.recommendation.recommended.params["engine"],
            second.recommendation.recommended.topic_count,
            second.recommendation.recommended.score,
            [(topic.count, topic.keywords) for topic in second.selected_topics],
        )
        self.assertEqual(first_signature, second_signature)

    def test_user_edit_is_recorded_and_export_state_changes(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "demo.xlsx"
            create_demo_workbook(workbook_path)
            package = analyze_file(workbook_path)
            topic_id = package.selected_topics[0].topic_id
            edit = apply_edit(
                package,
                "rename_topic",
                {"topic_id": topic_id, "label": "평가와 보상", "summary": "평가/보상 관련 개선 요구"},
                user_id="tester",
            )
            self.assertEqual(edit.edit_type, "rename_topic")
            self.assertEqual(package.selected_topics[0].label, "평가와 보상")
            self.assertEqual(len(package.user_edits), 1)
            approval = apply_edit(
                package,
                "record_review",
                {"decision": "approved", "notes": "대표 응답과 개인정보 확인"},
                user_id="reviewer-1",
            )
            self.assertEqual(approval.edit_type, "record_review")
            self.assertEqual(package.project["human_review"]["status"], "approved")
            apply_edit(
                package,
                "rename_topic",
                {"topic_id": topic_id, "label": "평가·보상 검토", "summary": "재검토"},
                user_id="tester",
            )
            self.assertEqual(package.project["human_review"]["status"], "changes_pending_review")

    def test_user_key_embedding_candidates_join_topic_count_competition(self) -> None:
        def fake_embeddings(texts: list[str], provider: str | None, api_key: str | None, **_: object) -> EmbeddingResult:
            vectors = np.asarray(
                [[1.0, 0.0, 0.0] if idx < len(texts) / 2 else [0.0, 1.0, 0.0] for idx, _ in enumerate(texts)],
                dtype=float,
            )
            return EmbeddingResult(
                vectors=vectors,
                source="openai_embeddings",
                model="fake-embedding-model",
                warnings=[],
            )

        with tempfile.TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "demo.xlsx"
            create_demo_workbook(workbook_path)
            with patch("survey_insight.topics.embed_texts_with_provider", side_effect=fake_embeddings):
                package = analyze_file(
                    workbook_path,
                    embedding_provider="openai",
                    embedding_api_key="sk-realistic-key-for-test-without-network",
                )

        engines = {candidate.params["engine"] for candidate in package.recommendation.candidates}
        self.assertIn("openai_embedding_kmeans", engines)
        embedding_candidates = [
            candidate
            for candidate in package.recommendation.candidates
            if candidate.params["engine"].startswith("openai_embedding")
        ]
        self.assertTrue(embedding_candidates)
        self.assertTrue(all(candidate.params["embedding_model"] == "fake-embedding-model" for candidate in embedding_candidates))


if __name__ == "__main__":
    unittest.main()
