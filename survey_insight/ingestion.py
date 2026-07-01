from __future__ import annotations

import csv
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models import TableData


def load_table(path: str | Path, sheet_name: str | None = None) -> TableData:
    source = Path(path)
    data = source.read_bytes()
    return load_table_bytes(data, source.name, sheet_name=sheet_name)


def load_table_bytes(
    data: bytes, filename: str, sheet_name: str | None = None
) -> TableData:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return _load_csv(data, filename)
    if suffix in {".xlsx", ".xlsm"}:
        return _load_xlsx(data, filename, sheet_name)
    raise ValueError(f"Unsupported file type: {suffix or filename}")


def _load_csv(data: bytes, filename: str) -> TableData:
    text = _decode_csv(data)
    rows = list(csv.reader(StringIO(text)))
    headers, records, header_idx = _rows_to_records(rows)
    return TableData(
        headers=headers,
        rows=records,
        sheet_name=None,
        header_row_index=header_idx,
        source_name=filename,
    )


def _load_xlsx(data: bytes, filename: str, sheet_name: str | None) -> TableData:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    raw_rows: list[list[Any]] = []
    for row in sheet.iter_rows(values_only=True):
        raw_rows.append(list(row))
    headers, records, header_idx = _rows_to_records(raw_rows)
    return TableData(
        headers=headers,
        rows=records,
        sheet_name=sheet.title,
        header_row_index=header_idx,
        source_name=filename,
    )


def _decode_csv(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _rows_to_records(raw_rows: list[list[Any]]) -> tuple[list[str], list[dict[str, Any]], int]:
    if not raw_rows:
        return [], [], 0
    header_idx = detect_header_row(raw_rows)
    header_cells = raw_rows[header_idx]
    width = max(len(row) for row in raw_rows)
    headers = [_normalize_header(cell, idx) for idx, cell in enumerate(_pad(header_cells, width))]
    records: list[dict[str, Any]] = []
    for raw in raw_rows[header_idx + 1 :]:
        padded = _pad(raw, len(headers))
        if all(_is_empty(cell) for cell in padded):
            continue
        records.append({headers[idx]: padded[idx] for idx in range(len(headers))})
    return headers, records, header_idx


def detect_header_row(raw_rows: list[list[Any]], max_scan_rows: int = 20) -> int:
    best_idx = 0
    best_score = float("-inf")
    width = max(len(row) for row in raw_rows)
    for idx, row in enumerate(raw_rows[:max_scan_rows]):
        padded = _pad(row, width)
        non_empty = [cell for cell in padded if not _is_empty(cell)]
        if not non_empty:
            continue
        strings = [cell for cell in non_empty if isinstance(cell, str)]
        unique = len({str(cell).strip() for cell in non_empty})
        next_non_empty = 0
        if idx + 1 < len(raw_rows):
            next_non_empty = sum(not _is_empty(cell) for cell in _pad(raw_rows[idx + 1], width))
        score = (
            len(non_empty) * 2.0
            + (unique / max(1, len(non_empty))) * 4.0
            + (len(strings) / max(1, len(non_empty))) * 3.0
            + min(next_non_empty, len(non_empty)) * 0.25
        )
        if any(str(cell).strip().lower() in {"name", "email", "응답", "의견"} for cell in non_empty):
            score += 1.0
        if score > best_score:
            best_idx = idx
            best_score = score
    return best_idx


def _normalize_header(cell: Any, idx: int) -> str:
    text = "" if cell is None else str(cell).strip()
    return text or f"column_{idx + 1}"


def _is_empty(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _pad(row: list[Any], width: int) -> list[Any]:
    if len(row) >= width:
        return row[:width]
    return row + [None] * (width - len(row))
