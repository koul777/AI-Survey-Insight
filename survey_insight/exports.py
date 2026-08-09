from __future__ import annotations

import html
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .models import AnalysisPackage


METRIC_LABELS = {
    "stability": "군집 품질 종합점수(휴리스틱)",
    "coherence": "토픽 해석 가능성 점수(UMass 방식)",
    "semantic_quality": "군집 분리도(cosine silhouette)",
    "coverage": "분석 포함률",
    "diversity": "토픽 키워드 다양성",
    "labelability": "라벨 해석 가능성",
    "balance": "토픽 크기 균형",
    "resampling_stability": "부분표본 일치도(진단)",
    "resampling_stability_p10": "부분표본 일치도 10백분위(진단)",
    "resampling_stability_p90": "부분표본 일치도 90백분위(진단)",
    "weight_acceptability": "가중치 시나리오 선정률(민감도)",
    "weight_acceptability_mcse": "가중치 시나리오 선정률 Monte Carlo 표준오차",
    "weight_score_p10": "가중치 시나리오 점수 10백분위",
    "weight_score_p90": "가중치 시나리오 점수 90백분위",
    "bootstrap_gate_pass_rate": "Bootstrap 구조 관문 통과율",
    "bootstrap_coverage_p025": "Bootstrap 분석 포함률 2.5백분위",
    "bootstrap_coverage_p975": "Bootstrap 분석 포함률 97.5백분위",
    "bootstrap_balance_p025": "Bootstrap 토픽 크기 균형 2.5백분위",
    "bootstrap_min_topic_share_p025": "Bootstrap 최소 토픽 비율 2.5백분위",
    "penalty": "품질 감점",
}

METHODOLOGY_SUMMARY = (
    "권장 토픽 수는 군집 분리도, 키워드 동시출현 기반 해석 가능성, 분석 포함률, "
    "키워드 다양성, 라벨 해석 가능성, 토픽 크기 균형을 결합한 휴리스틱으로 선정했습니다. "
    "가중치를 명시된 범위에서 바꾼 512개 시나리오의 선정률과 500회 assignment-conditional bootstrap을 "
    "별도 진단하고, 선정 후보는 5회 80% 층화 부분표본으로 다시 적합합니다. "
    "통계적으로 검증된 최적값이나 정확도가 아니므로 대표 응답과 함께 담당자가 검토해야 합니다."
)


def export_excel(package: AnalysisPackage, path: str | Path) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    ws = workbook.active
    ws.title = "원자료_토픽배정"
    _write_assigned_rows(ws, package)
    _write_topic_summary(workbook.create_sheet("토픽요약"), package)
    _write_representatives(workbook.create_sheet("대표의견"), package)
    _write_cross_analysis(workbook.create_sheet("교차분석"), package)
    _write_model_settings(workbook.create_sheet("모델설정"), package)
    _write_user_edits(workbook.create_sheet("편집이력"), package)
    for sheet in workbook.worksheets:
        _style_sheet(sheet)
    workbook.save(target)
    return target


def export_word(package: AnalysisPackage, path: str | Path) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document_xml = _document_xml(package)
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", _content_types_xml())
        zf.writestr("_rels/.rels", _rels_xml())
        zf.writestr("word/document.xml", document_xml)
        zf.writestr("word/styles.xml", _styles_xml())
        zf.writestr("word/_rels/document.xml.rels", _document_rels_xml())
    return target


def export_powerpoint(package: AnalysisPackage, path: str | Path) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.core_properties.title = f"{package.project.get('name', 'Survey Insight')} topic report"
    prs.core_properties.subject = "AI survey free-response topic modeling"
    prs.core_properties.comments = package.methodology_note
    _ppt_title_slide(prs, package)
    _ppt_summary_slide(prs, package)
    _ppt_topics_slide(prs, package)
    _ppt_representatives_slide(prs, package)
    _ppt_methodology_slide(prs, package)
    prs.save(target)
    return target


def _write_assigned_rows(ws: Any, package: AnalysisPackage) -> None:
    assignment_by_doc = {assignment.document_id: assignment.topic_id for assignment in package.assignments}
    topic_label = {topic.topic_id: topic.label for topic in package.selected_topics}
    headers = ["row_index", "redacted_text", "privacy_review_flags", "topic_id", "topic_label"] + package.group_columns
    ws.append(headers)
    for doc in package.documents:
        topic_id = assignment_by_doc.get(doc.id, "")
        ws.append(
            [
                doc.row_index,
                doc.redacted_text,
                ", ".join(doc.privacy_review_flags),
                topic_id,
                topic_label.get(topic_id, ""),
                *[doc.metadata.get(column) for column in package.group_columns],
            ]
        )


def _write_topic_summary(ws: Any, package: AnalysisPackage) -> None:
    ws.append([
        "topic_id",
        "label",
        "summary",
        "plain_language_summary",
        "suggested_action",
        "count",
        "share",
        "sentiment_label",
        "sentiment_plain_language",
        "sentiment_score",
        "urgency_score",
        "sentiment_evidence",
        "sentiment_method",
        "keywords",
    ])
    for topic in package.selected_topics:
        ws.append([
            topic.topic_id,
            topic.label,
            topic.summary,
            _optional_text(topic, "plain_language_summary"),
            _optional_text(topic, "suggested_action"),
            topic.count,
            topic.share,
            topic.sentiment_label,
            _optional_text(topic, "sentiment_plain_language"),
            topic.sentiment_score,
            topic.urgency_score,
            ", ".join(topic.sentiment_evidence),
            topic.sentiment_method,
            ", ".join(topic.keywords),
        ])


def _write_representatives(ws: Any, package: AnalysisPackage) -> None:
    ws.append(["topic_id", "topic_label", "rank", "representative_response"])
    for topic in package.selected_topics:
        for idx, response in enumerate(topic.representative_responses, start=1):
            ws.append([topic.topic_id, topic.label, idx, response])


def _write_cross_analysis(ws: Any, package: AnalysisPackage) -> None:
    ws.append(["group_column", "group_value", "topic_id", "topic_label", "display_count", "share", "suppressed"])
    for row in package.cross_analysis:
        ws.append(
            [
                row.get("group_column"),
                row.get("group_value"),
                row.get("topic_id"),
                row.get("topic_label"),
                row.get("display_count"),
                row.get("share"),
                row.get("suppressed"),
            ]
        )


def _write_model_settings(ws: Any, package: AnalysisPackage) -> None:
    rec = package.recommendation
    ws.append(["field", "value"])
    ws.append(["mode", rec.mode])
    ws.append(["valid_response_count", rec.valid_response_count])
    ws.append(["allowed_topic_range", f"{rec.allowed_topic_range[0]}-{rec.allowed_topic_range[1]}"])
    ws.append(["minimum_topic_size", rec.minimum_topic_size])
    ws.append(["recommended_topic_count", rec.recommended.topic_count])
    ws.append(["recommended_score (복합 평가값·휴리스틱)", rec.recommended.score])
    ws.append(["analysis_seed", package.project.get("analysis_seed", 42)])
    ws.append(["engine", rec.recommended.params.get("engine")])
    ws.append(["feature_space", rec.recommended.params.get("feature_space")])
    ws.append(["metric_profile", rec.recommended.params.get("metric_profile")])
    ws.append(["quality_gate_passed", rec.recommended.params.get("quality_gate_passed")])
    ws.append(["weight_sensitivity_profile", rec.recommended.params.get("weight_sensitivity_profile")])
    ws.append(["weight_sensitivity_scenarios", rec.recommended.params.get("weight_sensitivity_scenarios")])
    ws.append(["weight_multiplier_range", str(rec.recommended.params.get("weight_multiplier_range"))])
    ws.append(["bootstrap_profile", rec.recommended.params.get("bootstrap_profile")])
    ws.append(["bootstrap_repeats", rec.recommended.params.get("bootstrap_repeats")])
    ws.append(["resampling_repeats", rec.recommended.params.get("resampling_repeats")])
    review = _human_review(package)
    ws.append(["human_review.status", review.get("status")])
    ws.append(["human_review.reviewer", review.get("reviewer", "")])
    ws.append(["human_review.reviewed_at", review.get("reviewed_at", "")])
    ws.append(["privacy_review_flag_count", package.project.get("privacy_review_flag_count", 0)])
    for key, value in rec.recommended.metrics.items():
        ws.append([f"metric.{key} ({METRIC_LABELS.get(key, key)})", value])
    for interval in rec.recommended.params.get("bootstrap_topic_share_intervals", []):
        ws.append(
            [
                f"bootstrap.topic_share.{interval.get('topic_id')}",
                f"estimate={interval.get('estimate')}, p025={interval.get('p025')}, p975={interval.get('p975')}",
            ]
        )
    ws.append(["text_column", package.text_column])
    ws.append(["methodology_note", package.methodology_note])
    for label, value in _recommendation_explanations(rec):
        ws.append([label, value])
    for warning in rec.warnings:
        ws.append(["warning", warning])


def _write_user_edits(ws: Any, package: AnalysisPackage) -> None:
    ws.append(["timestamp", "user_id", "edit_type", "before", "after"])
    for edit in package.user_edits:
        ws.append([edit.timestamp, edit.user_id, edit.edit_type, str(edit.before), str(edit.after)])


def _style_sheet(ws: Any) -> None:
    header_fill = PatternFill("solid", fgColor="E8F0FE")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for idx, column_cells in enumerate(ws.columns, start=1):
        width = min(60, max(10, max(len(str(cell.value or "")) for cell in column_cells[:50]) + 2))
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"


def _ppt_title_slide(prs: Presentation, package: AnalysisPackage) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = str(package.project.get("name", "Survey Insight Project"))
    subtitle = slide.placeholders[1]
    subtitle.text = (
        "AI 설문 자유응답 토픽모델링 분석\n"
        f"{package.recommendation.mode} · 유효 응답 {package.recommendation.valid_response_count}건"
    )


def _ppt_summary_slide(prs: Presentation, package: AnalysisPackage) -> None:
    slide = _ppt_blank_slide(prs, "핵심 발견")
    rec = package.recommendation.recommended
    lines = [
        f"총 응답 {package.dataset_profile.row_count}건 중 유효 자유응답 {package.recommendation.valid_response_count}건을 분석했습니다.",
        f"권장 토픽 수는 {rec.topic_count}개이며 복합 평가값(휴리스틱)은 {rec.score:.2f}입니다.",
        f"분석 컬럼: {package.text_column}",
        f"사람 검토 상태: {_human_review(package).get('status', 'not_reviewed')}",
    ]
    lines.extend(f"{label}: {value}" for label, value in _recommendation_explanations(package.recommendation))
    lines.extend(package.recommendation.warnings[:2])
    _ppt_bullets(slide, lines, left=0.8, top=1.4, width=8.8, height=3.4)
    _ppt_note(slide, package.methodology_note)


def _ppt_topics_slide(prs: Presentation, package: AnalysisPackage) -> None:
    slide = _ppt_blank_slide(prs, "주요 토픽")
    rows = min(len(package.selected_topics), 6) + 1
    table = slide.shapes.add_table(rows, 6, Inches(0.25), Inches(1.25), Inches(9.5), Inches(4.4)).table
    headers = ["토픽", "응답 수", "비율", "감정 신호/처리 우선도", "쉽게 읽는 해석", "대응 방안"]
    for idx, header in enumerate(headers):
        table.cell(0, idx).text = header
        _ppt_cell_header(table.cell(0, idx))
    for row_idx, topic in enumerate(package.selected_topics[:6], start=1):
        values = [
            f"{topic.topic_id} {topic.label}",
            str(topic.count),
            f"{topic.share:.1%}",
            f"{topic.sentiment_label} / {topic.urgency_score:.2f}",
            _optional_text(topic, "plain_language_summary") or topic.summary,
            _optional_text(topic, "suggested_action") or ", ".join(topic.keywords[:5]),
        ]
        for col_idx, value in enumerate(values):
            table.cell(row_idx, col_idx).text = value
            _ppt_cell_body(table.cell(row_idx, col_idx), center=col_idx in {1, 2, 3})


def _ppt_representatives_slide(prs: Presentation, package: AnalysisPackage) -> None:
    slide = _ppt_blank_slide(prs, "대표 의견")
    bullets: list[str] = []
    for topic in package.selected_topics[:4]:
        if topic.representative_responses:
            bullets.append(f"{topic.label}: {topic.representative_responses[0]}")
    if not bullets:
        bullets.append("대표 의견을 생성할 유효 응답이 없습니다.")
    _ppt_bullets(slide, bullets, left=0.65, top=1.15, width=9.0, height=4.8)


def _ppt_methodology_slide(prs: Presentation, package: AnalysisPackage) -> None:
    slide = _ppt_blank_slide(prs, "방법론과 주의사항")
    rec = package.recommendation.recommended
    lines = [
        METHODOLOGY_SUMMARY,
        f"사용 엔진: {rec.params.get('engine', 'unknown')}",
        f"허용 토픽 범위: {package.recommendation.allowed_topic_range[0]}-{package.recommendation.allowed_topic_range[1]}개",
        "표시 지표: 군집 품질 종합점수, 토픽 해석 가능성 점수, 군집 분리도, 분석 포함률, 토픽 키워드 다양성, 라벨 해석 가능성",
        "가중치 시나리오 선정률은 512개 제한 범위 가중치 조합에서 같은 후보가 선택된 비율이며 정확도 확률이 아닙니다.",
        "Bootstrap 구조 관문 통과율은 현재 배정에 조건부인 500회 진단이며 모집단 타당성이나 검정력이 아닙니다.",
        "부분표본 일치도는 선정 후보의 5회 재적합 사후 진단이며 권장 점수에는 포함되지 않습니다.",
        "감정은 규칙·어휘 기반 신호이며 처리 우선도는 긴급성 판정이 아닙니다.",
        package.methodology_note,
    ]
    _ppt_bullets(slide, lines, left=0.65, top=1.2, width=9.0, height=4.5)


def _ppt_blank_slide(prs: Presentation, title: str) -> Any:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)
    slide.shapes.title.text_frame.paragraphs[0].font.bold = True
    return slide


def _ppt_bullets(slide: Any, lines: list[str], left: float, top: float, width: float, height: float) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for idx, line in enumerate(lines):
        paragraph = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        paragraph.text = line
        paragraph.level = 0
        paragraph.font.size = Pt(16)
        paragraph.space_after = Pt(8)


def _ppt_note(slide: Any, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.65), Inches(6.45), Inches(9.0), Inches(0.35))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(9)
    p.font.italic = True


def _ppt_cell_header(cell: Any) -> None:
    for paragraph in cell.text_frame.paragraphs:
        paragraph.font.bold = True
        paragraph.font.size = Pt(12)
        paragraph.alignment = PP_ALIGN.CENTER


def _ppt_cell_body(cell: Any, center: bool = False) -> None:
    for paragraph in cell.text_frame.paragraphs:
        paragraph.font.size = Pt(10)
        paragraph.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT


def _document_xml(package: AnalysisPackage) -> str:
    body: list[str] = []
    project_name = package.project.get("name", "Survey Insight Project")
    body.append(_p(project_name, style="Title"))
    body.append(_p("AI 설문 자유응답 토픽모델링 분석 보고서", style="Subtitle"))
    body.append(_p(package.methodology_note))
    body.append(_p(f"사람 검토 상태: {_human_review(package).get('status', 'not_reviewed')}"))
    for warning in package.recommendation.warnings:
        body.append(_p(f"주의: {warning}", style="Warning"))
    body.append(_heading("Executive Summary", 1))
    body.append(
        _p(
            f"총 {package.dataset_profile.row_count}개 응답 중 "
            f"{package.recommendation.valid_response_count}개의 유효 자유응답을 분석했습니다. "
            f"권장 토픽 수는 {package.recommendation.recommended.topic_count}개입니다."
        )
    )
    for label, value in _recommendation_explanations(package.recommendation):
        body.append(_heading(label, 2))
        body.append(_p(value))
    body.append(_heading("권장 근거", 1))
    rec = package.recommendation.recommended
    metric_rows = [["표시 지표", "내부 필드", "점수"]] + [
        [METRIC_LABELS.get(key, key), key, value] for key, value in rec.metrics.items()
    ]
    body.append(_table(metric_rows))
    body.append(_heading("주요 토픽", 1))
    topic_rows = [["토픽", "응답 수", "비율", "감정 신호", "처리 우선도", "감정 방식", "키워드", "요약", "쉽게 읽는 해석", "대응 방안", "감정 해석"]]
    for topic in package.selected_topics:
        topic_rows.append(
            [
                f"{topic.topic_id} {topic.label}",
                str(topic.count),
                f"{topic.share:.1%}",
                f"{topic.sentiment_label} ({topic.sentiment_score:+.2f})",
                f"{topic.urgency_score:.2f}",
                topic.sentiment_method,
                ", ".join(topic.keywords),
                topic.summary,
                _optional_text(topic, "plain_language_summary"),
                _optional_text(topic, "suggested_action"),
                _optional_text(topic, "sentiment_plain_language"),
            ]
        )
    body.append(_table(topic_rows))
    body.append(_heading("대표 의견", 1))
    for topic in package.selected_topics:
        body.append(_heading(f"{topic.topic_id} {topic.label}", 2))
        for label, value in _topic_plain_language_sections(topic):
            body.append(_p(f"{label}: {value}"))
        for response in topic.representative_responses:
            body.append(_p(response, style="Quote"))
    if package.cross_analysis:
        body.append(_heading("응답자 특성별 인사이트", 1))
        rows = [["그룹", "값", "토픽", "표시 건수", "비율"]]
        for row in package.cross_analysis[:40]:
            rows.append(
                [
                    str(row.get("group_column")),
                    str(row.get("group_value")),
                    str(row.get("topic_label")),
                    "숨김" if row.get("suppressed") else str(row.get("display_count")),
                    f"{float(row.get('share') or 0):.1%}",
                ]
            )
        body.append(_table(rows))
    body.append(_heading("방법론 부록", 1))
    body.append(_p(f"분석 컬럼: {package.text_column}"))
    body.append(_p(f"분석 모드: {package.recommendation.mode}"))
    body.append(_p(METHODOLOGY_SUMMARY))
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body>"
        + "".join(body)
        + '<w:sectPr><w:pgSz w:w="12240" w:h="15840"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/></w:sectPr>'
        "</w:body></w:document>"
    )


def _optional_text(obj: Any, field_name: str) -> str:
    value = getattr(obj, field_name, "")
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return " ".join(str(item) for item in value if item)
    return str(value)


def _human_review(package: AnalysisPackage) -> dict[str, Any]:
    review = package.project.get("human_review")
    if isinstance(review, dict):
        return review
    return {"status": "not_reviewed"}


def _recommendation_explanations(rec: Any) -> list[tuple[str, str]]:
    recommended = getattr(rec, "recommended", None)
    entries = [
        (
            "쉽게 읽는 해석",
            _first_optional_text(rec, ["plain_language_summary", "summary_plain_language"])
            or _first_optional_text(recommended, ["plain_language_summary", "summary_plain_language"]),
        ),
        (
            "왜 이 토픽 수를 권장했나요?",
            _first_optional_text(rec, ["topic_count_explanation", "recommendation_reason", "topic_count_reason", "why_recommended"])
            or _first_optional_text(recommended, ["recommendation_reason", "topic_count_reason", "why_recommended"]),
        ),
        (
            "주의해서 읽을 점",
            _first_optional_text(rec, ["quality_warnings", "reading_cautions", "plain_language_caution", "caution"]),
        ),
        (
            "방법론 한 줄 설명",
            _first_optional_text(rec, ["methodology_plain_language"]),
        ),
    ]
    return [(label, value) for label, value in entries if value]


def _topic_plain_language_sections(topic: Any) -> list[tuple[str, str]]:
    entries = [
        ("쉽게 읽는 해석", _optional_text(topic, "plain_language_summary")),
        ("대응 방안", _optional_text(topic, "suggested_action")),
        ("감정 해석", _optional_text(topic, "sentiment_plain_language")),
    ]
    return [(label, value) for label, value in entries if value]


def _first_optional_text(obj: Any, field_names: list[str]) -> str:
    if obj is None:
        return ""
    for field_name in field_names:
        value = _optional_text(obj, field_name)
        if value:
            return value
    return ""


def _p(text: Any, style: str | None = None) -> str:
    escaped = _escape(text)
    style_xml = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
    return f"<w:p>{style_xml}<w:r><w:t>{escaped}</w:t></w:r></w:p>"


def _heading(text: str, level: int) -> str:
    return _p(text, style=f"Heading{level}")


def _table(rows: list[list[Any]]) -> str:
    table_rows = []
    for row_idx, row in enumerate(rows):
        cells = []
        for value in row:
            style = '<w:pPr><w:pStyle w:val="TableHeader"/></w:pPr>' if row_idx == 0 else ""
            cells.append(
                "<w:tc><w:tcPr><w:tcW w:w=\"2400\" w:type=\"dxa\"/></w:tcPr>"
                f"<w:p>{style}<w:r><w:t>{_escape(value)}</w:t></w:r></w:p></w:tc>"
            )
        table_rows.append("<w:tr>" + "".join(cells) + "</w:tr>")
    return (
        '<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/><w:tblBorders>'
        '<w:top w:val="single" w:sz="4" w:color="D0D7DE"/>'
        '<w:left w:val="single" w:sz="4" w:color="D0D7DE"/>'
        '<w:bottom w:val="single" w:sz="4" w:color="D0D7DE"/>'
        '<w:right w:val="single" w:sz="4" w:color="D0D7DE"/>'
        '<w:insideH w:val="single" w:sz="4" w:color="D0D7DE"/>'
        '<w:insideV w:val="single" w:sz="4" w:color="D0D7DE"/>'
        "</w:tblBorders></w:tblPr>"
        + "".join(table_rows)
        + "</w:tbl>"
    )


def _escape(value: Any) -> str:
    return html.escape("" if value is None else str(value), quote=False)


def _content_types_xml() -> str:
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
</Types>"""


def _rels_xml() -> str:
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""


def _document_rels_xml() -> str:
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>"""


def _styles_xml() -> str:
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:eastAsia="Malgun Gothic"/><w:sz w:val="22"/></w:rPr></w:style>
  <w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/><w:rPr><w:b/><w:sz w:val="36"/></w:rPr></w:style>
  <w:style w:type="paragraph" w:styleId="Subtitle"><w:name w:val="Subtitle"/><w:basedOn w:val="Normal"/><w:rPr><w:color w:val="555555"/><w:sz w:val="24"/></w:rPr></w:style>
  <w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/><w:pPr><w:spacing w:before="320" w:after="120"/></w:pPr><w:rPr><w:b/><w:sz w:val="28"/></w:rPr></w:style>
  <w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/><w:basedOn w:val="Normal"/><w:pPr><w:spacing w:before="220" w:after="80"/></w:pPr><w:rPr><w:b/><w:sz w:val="24"/></w:rPr></w:style>
  <w:style w:type="paragraph" w:styleId="Warning"><w:name w:val="Warning"/><w:basedOn w:val="Normal"/><w:rPr><w:b/><w:color w:val="9A3412"/></w:rPr></w:style>
  <w:style w:type="paragraph" w:styleId="Quote"><w:name w:val="Quote"/><w:basedOn w:val="Normal"/><w:pPr><w:ind w:left="360"/></w:pPr><w:rPr><w:color w:val="374151"/></w:rPr></w:style>
  <w:style w:type="paragraph" w:styleId="TableHeader"><w:name w:val="Table Header"/><w:basedOn w:val="Normal"/><w:rPr><w:b/></w:rPr></w:style>
</w:styles>"""
